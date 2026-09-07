"""Render the versioned FoodLens content source into an editable official draft.

Run with Codex's bundled python-docx runtime. Export this DOCX via the packaged
render_docx.py, then verify using check-competition-proposal.py. Never fill real
identities or imply submission approval in a generated draft.
"""

import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/competition/foodlens-proposal-v0.1.json"
CONTENT = json.loads(SOURCE.read_text())
STEM = "FoodLens_初選提案_官方格式待填稿_v0.1"
OUT = ROOT / "output/competition"
WORK = OUT / "work/foodlens-v01"
FONT = CONTENT["font"]
DOC = Document()
LEDGER = []
ASSETS = []
PAGE = 1


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def set_font(run, size=12, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    prop = run._element.get_or_add_rPr()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        prop.rFonts.set(qn("w:" + key), FONT)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "zh-TW")
    lang.set(qn("w:eastAsia"), "zh-TW")
    prop.append(lang)


def record(text, kind="paragraph"):
    if text.strip():
        LEDGER.append({"page": PAGE, "kind": kind, "text": text})


def para(text, style=None, before=0, after=7, size=12, bold=False, align=None):
    p = DOC.add_paragraph(style=style)
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = Pt(18)
    pf.widow_control = True
    if align is not None:
        p.alignment = align
    set_font(p.add_run(text), size, bold)
    record(text)
    return p


def heading(text, level=1):
    p = para(text, "Heading " + str(level), before=0 if level == 1 else 6,
             after=10 if level == 1 else 7, size=14, bold=True)
    p.paragraph_format.line_spacing = Pt(22)
    p.paragraph_format.keep_with_next = True
    return p


def new_page(number, title):
    global PAGE
    DOC.add_page_break()
    PAGE = number
    heading(title)


def add_link(p, label, url):
    hyperlink = OxmlElement("w:hyperlink")
    rid = p.part.relate_to(url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    hyperlink.set(qn("r:id"), rid)
    run = p.add_run(label)
    set_font(run)
    run.font.underline = True
    hyperlink.append(run._element)
    p._p.append(hyperlink)
    record(label, "hyperlink")


def table(block):
    headers, rows, widths = block["headers"], block["rows"], block["widths"]
    assert sum(widths) == 165
    t = DOC.add_table(rows=1, cols=len(headers))
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for column, width in zip(t.columns, widths):
        column.width = Mm(width)
    pr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement("w:" + edge)
        for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
            b.set(qn("w:" + key), value)
        borders.append(b)
    pr.append(borders)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i, values in enumerate([headers] + rows):
        row = t.rows[0] if i == 0 else t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, width, value in zip(row.cells, widths, values):
            cell.width = Mm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            mar = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                m = OxmlElement("w:" + edge)
                m.set(qn("w:w"), "85")
                m.set(qn("w:type"), "dxa")
                mar.append(m)
            tcpr.append(mar)
            if i == 0:
                fill = OxmlElement("w:shd")
                fill.set(qn("w:fill"), "EEEEEE")
                tcpr.append(fill)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(17)
            set_font(p.add_run(value), bold=i == 0)
            record(value, "table-cell")
    tail = DOC.add_paragraph()
    tail.paragraph_format.space_after = Pt(3)
    tail.paragraph_format.line_spacing = Pt(1)
    set_font(tail.add_run(""))


def picture(block):
    source = ROOT / block["source"]
    image = Image.open(source)
    crop = tuple(block["crop"])
    assert 0 <= crop[0] < crop[2] <= image.width
    assert 0 <= crop[1] < crop[3] <= image.height
    cropped = WORK / (source.stem + "-document-crop.png")
    image.crop(crop).save(cropped)
    p = DOC.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    pic = p.add_run().add_picture(str(cropped), width=Mm(165))
    pic._inline.docPr.set("descr", block["alt"])
    pic._inline.docPr.set("title", "FoodLens 原型截圖（示範資料）")
    para(block["caption"], after=7)
    ASSETS.append({"source": block["source"], "sha256": sha(source),
                   "crop": list(crop), "crop_sha256": sha(cropped),
                   "page": PAGE, "alt": block["alt"], "caption": block["caption"]})


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    WORK.mkdir(parents=True, exist_ok=True)
    assert len(CONTENT["abstract"]) <= 300
    assert len(CONTENT["abstract_short"]) <= 100
    assert [p["number"] for p in CONTENT["pages"]] == list(range(3, 11))
    sec = DOC.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(22.5)
    sec.top_margin = Mm(20)
    sec.bottom_margin = Mm(19)
    sec.header_distance = sec.footer_distance = Mm(9)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        style = DOC.styles[name]
        style.font.name = FONT
        style.font.size = Pt(12 if name in ("Normal", "Subtitle") else 14)
        style.font.color.rgb = RGBColor(0, 0, 0)
        # The bundled blank DOCX may carry a blue Title border. Official
        # submissions use the supplied plain template, not that decoration.
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
        rpr = style.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.insert(0, rf)
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            rf.set(qn("w:" + key), FONT)
    DOC.styles["Title"].font.size = Pt(24)
    default_lang = OxmlElement("w:themeFontLang")
    default_lang.set(qn("w:val"), "zh-TW")
    default_lang.set(qn("w:eastAsia"), "zh-TW")
    DOC.settings.element.append(default_lang)
    DOC.core_properties.title = CONTENT["title"] + "｜初選提案構想書"
    DOC.core_properties.subject = CONTENT["category"]
    DOC.core_properties.author = "FoodLens 專案；參賽學生待具名覆核"
    DOC.core_properties.language = "zh-TW"
    DOC.core_properties.comments = CONTENT["status"]
    DOC.core_properties.created = datetime(2026, 9, 6, tzinfo=timezone.utc)
    DOC.core_properties.modified = datetime(2026, 9, 6, tzinfo=timezone.utc)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("FoodLens　初選待填稿 v0.1　｜　"))
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)

    para("臺北市115學年度", align=WD_ALIGN_PARAGRAPH.CENTER, before=22, size=14)
    para("智慧城市中的科技與人文：\n跨域創新素養競賽", align=WD_ALIGN_PARAGRAPH.CENTER, after=74, size=14)
    p = para(CONTENT["title"], "Title", align=WD_ALIGN_PARAGRAPH.CENTER, size=24, bold=True, after=12)
    p.paragraph_format.line_spacing = Pt(32)
    para(CONTENT["subtitle"], "Subtitle", align=WD_ALIGN_PARAGRAPH.CENTER, after=52, size=14)
    para("參選類別：" + CONTENT["category"], after=12)
    for item in ("學層組別：國小組", "團隊名稱：待填", "參選人姓名：待填（2至4名）", "學校名稱：待填", "指導教師：待確認（可選擇1名）"):
        para(item, after=12)
    para("官方格式待填稿 v0.1｜2026年9月6日", before=28)
    para("本稿尚未送件；參賽身分、學生實際貢獻與內容素材仍待具名覆核。", after=0)

    new_page(2, "內容簡介")
    para(CONTENT["abstract"], before=12, after=18)
    for page in CONTENT["pages"]:
        new_page(page["number"], page["title"])
        for block in page["blocks"]:
            if block["type"] == "heading":
                heading(block["text"], level=2)
            elif block["type"] == "paragraph":
                p = para(block["text"])
                if "link" in block:
                    add_link(p, block["link_label"], block["link"])
            elif block["type"] == "table":
                table(block)
            elif block["type"] == "image":
                picture(block)
            else:
                raise ValueError(block["type"])
    path = OUT / (STEM + ".docx")
    DOC.save(path)
    ledger = {"schema_version": 1, "content_sha256": sha(SOURCE),
              "builder_sha256": sha(Path(__file__)), "docx_sha256": sha(path),
              "abstract_characters": len(CONTENT["abstract"]),
              "entry_characters": len(CONTENT["abstract_short"]),
              "snapshot": CONTENT["snapshot"], "assets": ASSETS,
              "paragraphs": LEDGER, "validation": "pending DOCX-to-PDF render and visual review"}
    (WORK / "source-ledger.json").write_text(json.dumps(ledger, ensure_ascii=False, indent=2) + "\n")
    md = ["# FoodLens 官方格式待填稿 v0.1（產生稿）", "",
          "唯一內容來源：docs/competition/foodlens-proposal-v0.1.json。請先修改來源，再重建 DOCX；PDF 從該 DOCX 轉出。", "",
          "## 報名提案理念（" + str(len(CONTENT["abstract_short"])) + " 字）", "", CONTENT["abstract_short"], ""]
    prior = None
    for item in LEDGER:
        if item["page"] != prior:
            md += ["## 第 " + str(item["page"]) + " 頁", ""]
            prior = item["page"]
        md += [item["text"], ""]
    (OUT / (STEM + ".md")).write_text("\n".join(md))
    print(json.dumps({"docx": str(path), "paragraphs": len(LEDGER),
                      "abstract_chars": len(CONTENT["abstract"]), "entry_chars": len(CONTENT["abstract_short"])}))


if __name__ == "__main__":
    build()
