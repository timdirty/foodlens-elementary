"""Native three-page field record; blank paper values never become Demo data."""

import argparse
from datetime import datetime, timezone
import hashlib
import importlib.util
import json
from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/competition/field-record-v0.1.json"
STEM = "FoodLens_單餐量測紀錄表_v0.1"
DOCX = ROOT / "output/competition" / (STEM + ".docx")
WORK = ROOT / "output/competition/work/field-record-v01"
FONT = "BiauKaiTC"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def font(run, size=12, bold=False):
    run.font.name, run.font.size = FONT, Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    props = run._element.get_or_add_rPr()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        props.rFonts.set(qn("w:" + key), FONT)
    language = OxmlElement("w:lang")
    language.set(qn("w:val"), "zh-TW")
    language.set(qn("w:eastAsia"), "zh-TW")
    props.append(language)


def build():
    content = json.loads(SOURCE.read_text(encoding="utf-8"))
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(18)
    section.top_margin, section.bottom_margin = Mm(15), Mm(15)
    section.footer_distance = Mm(7)
    for name, size in (("Normal", 12), ("Title", 24), ("Heading 1", 14)):
        style = doc.styles[name]
        style.font.name, style.font.size = FONT, Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
        pr = style.element.get_or_add_rPr()
        rf = pr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            pr.insert(0, rf)
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            rf.set(qn("w:" + key), FONT)
    doc.core_properties.title = content["title"]
    doc.core_properties.author = "FoodLens 專案"
    doc.core_properties.subject = content["status"]
    doc.core_properties.language = "zh-TW"
    doc.core_properties.created = datetime(2026, 9, 6, tzinfo=timezone.utc)
    doc.core_properties.modified = datetime(2026, 9, 6, tzinfo=timezone.utc)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run("FoodLens　空白量測表 v0.1　非報名附件　｜　"))
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    ledger = []
    for number, page in enumerate(content["pages"], 1):
        if number > 1:
            doc.add_page_break()
        p = doc.add_paragraph(style="Title")
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = Pt(30)
        font(p.add_run(page["title"]), 24, True)
        ledger.append({"page": number, "text": page["title"], "role": "title"})
        for block in page["blocks"]:
            kind = block["type"]
            if kind != "table":
                p = doc.add_paragraph(style="Heading 1" if kind == "heading" else "Normal")
                pf = p.paragraph_format
                pf.space_before = Pt(6 if kind == "heading" else 0)
                pf.space_after = Pt(block.get("after", 5 if kind == "heading" else 4))
                pf.line_spacing = Pt(20 if kind == "heading" else 17)
                pf.widow_control = True
                pf.keep_with_next = kind == "heading"
                font(p.add_run(block["text"]), 14 if kind == "heading" else 12, kind == "heading")
                ledger.append({"page": number, "text": block["text"], "role": kind})
                continue
            widths = block["widths"]
            assert sum(widths) == 174
            table = doc.add_table(rows=1, cols=len(widths))
            table.autofit = False
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for column, width in zip(table.columns, widths):
                column.width = Mm(width)
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                element = OxmlElement("w:" + edge)
                for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    element.set(qn("w:" + key), value)
                borders.append(element)
            table._tbl.tblPr.append(borders)
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for index, values in enumerate([block["headers"]] + block["rows"]):
                row = table.rows[0] if index == 0 else table.add_row()
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                if index > 0:
                    row.height, row.height_rule = Mm(block["minimum_row_mm"]), WD_ROW_HEIGHT_RULE.AT_LEAST
                for cell, width, value in zip(row.cells, widths, values):
                    cell.width = Mm(width)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    pr = cell._tc.get_or_add_tcPr()
                    margins = OxmlElement("w:tcMar")
                    for edge in ("top", "left", "bottom", "right"):
                        m = OxmlElement("w:" + edge)
                        m.set(qn("w:w"), "70")
                        m.set(qn("w:type"), "dxa")
                        margins.append(m)
                    pr.append(margins)
                    if index == 0:
                        fill = OxmlElement("w:shd")
                        fill.set(qn("w:fill"), "EEEEEE")
                        pr.append(fill)
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(widths) >= 4 else WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = Pt(16)
                    font(p.add_run(value), bold=index == 0)
                    if value:
                        ledger.append({"page": number, "text": value, "role": "table-cell"})
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = Pt(1)
    DOCX.parent.mkdir(parents=True, exist_ok=True)
    WORK.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX)
    (WORK / "source-ledger.json").write_text(json.dumps({
        "source_sha256": sha(SOURCE), "builder_sha256": sha(Path(__file__)),
        "docx_sha256": sha(DOCX), "expected_pages": 3, "font": FONT,
        "paragraphs": ledger,
        "sources": [{"path": r, "sha256": sha(ROOT / r)} for r in content["source_refs"]],
        "status": "blank-form-awaiting-render-and-visual-check",
        "interpretation": "Blank editable Word text and print PDF; not AcroForm, not collected data or school authorization.",
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(str(DOCX))


def render(renderer):
    spec = importlib.util.spec_from_file_location("field_record_renderer", ROOT / "scripts/render-competition-proposal.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    module.STEM, module.WORK = STEM, WORK
    sys.argv = [str(Path(__file__)), "--renderer", str(renderer)]
    module.main()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--renderer", type=Path, help="Optional packaged render_docx.py; uses the same bundled Python and task-local font setup")
    args = parser.parse_args()
    build()
    if args.renderer:
        render(args.renderer)
